from io import BytesIO
import zipfile

from docx import Document
import fitz
from openpyxl import load_workbook
import pytest

from app.reporting import render_report


ROWS = [{"created_at":"2026-10-20T12:00:00+00:00","domain":"bcp-control.invalid","title":"=HYPERLINK(\"https://example.com\")","score":65,"risk":"high","severity":"high","status":"open","source":"manual","message":"Posible suplantación; <script>no ejecutar</script>","timeline":[{"at":"2026-10-20T12:00:00Z","action":"reviewed","note":"Verificado por canal oficial"}],"record_type":"evaluations"}]


@pytest.mark.parametrize("kind",["evaluations","alerts","incidents","consolidated"])
@pytest.mark.parametrize("format",["pdf","docx","xlsx"])
def test_real_formats_contain_authorized_data_and_metadata(kind,format):
    payload=render_report(kind,format,ROWS,{"requester":"Analista de pruebas","filters":{"risk":"high"}})
    assert len(payload)>1000
    if format=="pdf":
        document=fitz.open(stream=payload,filetype="pdf")
        text="".join(page.get_text() for page in document)
        assert "Analista de pruebas" in text and "Página 1" in text
        assert "alto" in text.lower()
    elif format=="docx":
        document=Document(BytesIO(payload))
        text="\n".join(p.text for p in document.paragraphs)
        assert "Analista de pruebas" in text
        assert document.tables and "Fecha UTC" in str([c.text for c in document.tables[0].rows[0].cells])
        assert "word/document.xml" in zipfile.ZipFile(BytesIO(payload)).namelist()
    else:
        workbook=load_workbook(BytesIO(payload))
        sheet=workbook["Eventos"]
        assert sheet.freeze_panes=="A2" and sheet.auto_filter.ref
        assert not any(cell.data_type=="f" for row in sheet for cell in row)
        assert workbook["Resumen"]["B4"].value=="Analista de pruebas"


def test_empty_report_is_explicit():
    document=fitz.open(stream=render_report("evaluations","pdf",[],{}),filetype="pdf")
    assert "No hay registros" in "".join(page.get_text() for page in document)


def test_long_incident_notes_split_across_pages_instead_of_failing():
    row={**ROWS[0],"timeline":[{"at":"2026-10-21T00:00:00Z","action":"reviewed","note":"evidencia "+"detalle "*300} for _ in range(3)]}
    payload=render_report("incidents","pdf",[row],{"requester":"Analista de pruebas"})
    document=fitz.open(stream=payload,filetype="pdf")
    assert len(document)>1
    text="".join(page.get_text() for page in document)
    assert "texto abreviado" in text


def test_dates_with_offset_are_converted_before_being_labeled_utc():
    rows=[{**ROWS[0],"created_at":"2026-10-21T14:00:00-05:00"}]
    workbook=load_workbook(BytesIO(render_report("evaluations","xlsx",rows,{})))
    assert workbook["Eventos"]["A2"].value=="2026-10-21 19:00:00 UTC"
