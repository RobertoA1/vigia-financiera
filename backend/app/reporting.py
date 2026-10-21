"""Data-only report rendering. Authorization and filtering happen in the API."""
from __future__ import annotations

from collections import Counter
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import json
import textwrap
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table as PDFTable, TableStyle, KeepTogether
from reportlab.graphics.shapes import Drawing, Rect, String
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

FONTS = Path(__file__).resolve().parents[1] / "assets" / "fonts"
pdfmetrics.registerFont(TTFont("Vigia", str(FONTS / "DejaVuSans.ttf")))
pdfmetrics.registerFont(TTFont("VigiaBold", str(FONTS / "DejaVuSans-Bold.ttf")))
pdfmetrics.registerFontFamily("Vigia", normal="Vigia", bold="VigiaBold", italic="Vigia", boldItalic="VigiaBold")

TITLES = {"evaluations": "Reporte de evaluaciones de URLs", "alerts": "Reporte de alertas", "incidents": "Reporte de incidentes", "consolidated": "Reporte operativo consolidado"}
RISK_LABELS = {"low": "Bajo", "moderate": "Moderado", "high": "Alto", "critical": "Crítico"}
STATUS_LABELS = {"preliminary": "Preliminar", "complete": "Completo", "pending": "Pendiente", "inconclusive": "No concluyente", "open": "Abierto", "investigating": "En investigación", "resolved": "Resuelto", "dismissed": "Descartado", "read": "Leída", "unread": "No leída"}
COLUMNS = {
    "evaluations": [("created_at", "Fecha UTC"), ("domain", "Dominio"), ("score", "Puntuación"), ("risk", "Riesgo"), ("status", "Estado"), ("source", "Origen")],
    "alerts": [("created_at", "Fecha UTC"), ("title", "Alerta"), ("risk", "Riesgo"), ("message", "Motivos"), ("read", "Leída"), ("notification_status", "Notificación")],
    "incidents": [("created_at", "Fecha UTC"), ("title", "Incidente"), ("severity", "Severidad"), ("status", "Estado"), ("assigned_to", "Responsable"), ("timeline", "Acciones y observaciones UTC")],
    "consolidated": [("record_type", "Categoría"), ("created_at", "Fecha UTC"), ("event_title", "Dominio o evento"), ("event_risk", "Puntuación y riesgo"), ("status", "Estado"), ("event_detail", "Evidencia y acciones")],
}


def _text(value) -> str:
    if value is None:
        return "Sin dato"
    if isinstance(value, datetime):
        return value.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, bool):
        return "Sí" if value else "No"
    value = str(value)
    clean = "".join(c for c in value if ord(c) >= 32 or c in "\n\t")
    return clean[:3900] + " [texto abreviado por extensión]" if len(clean) > 4000 else clean


def _values(row, columns):
    result = []
    for key, _ in columns:
        value = row.get(key)
        if key in ("risk", "severity"):
            value = RISK_LABELS.get(value, value)
        elif key == "status":
            value = STATUS_LABELS.get(value, value)
        elif key in {"timeline", "event_detail"} and isinstance(value, list):
            names = {"created":"Creado", "risk_updated":"Riesgo actualizado", "escalated":"Escalado", "updated":"Actualizado", "note":"Observación"}
            value = "\n".join(f"{event.get('at', '')} {names.get(event.get('action'), event.get('action', 'Acción'))}: {event.get('note', '')}" for event in value)
        elif key == "record_type":
            value = {"evaluations": "Evaluación", "alerts": "Alerta", "incidents": "Incidente"}.get(value, value)
        elif key == "source":
            value = {"manual": "Manual", "extension": "Extensión", "email": "Correo", "sms": "SMS", "api": "API"}.get(value, value)
        if key == "created_at" or key == "updated_at":
            if isinstance(value, str):
                try:
                    parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
                    value = parsed.replace(tzinfo=timezone.utc).astimezone(timezone.utc) if parsed.tzinfo is None else parsed.astimezone(timezone.utc)
                    value = value.strftime("%Y-%m-%d %H:%M:%S UTC")
                except ValueError:
                    value = "Fecha no disponible"
        result.append(value if isinstance(value, (int, float)) and not isinstance(value, bool) else _text(value))
    return result


def _metadata(kind, rows, metadata):
    filters = metadata.get("filters", {})
    requester = metadata.get("requested_by", metadata.get("requester", metadata.get("user", "Usuario autorizado")))
    if isinstance(requester, dict):
        requester = requester.get("email", requester.get("name", "Usuario autorizado"))
    return {
        "Sistema": metadata.get("system", metadata.get("brand", "Vigía Financiera")),
        "Generado UTC": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        "Solicitante": _text(requester),
        "Período": f"{filters.get('date_from', metadata.get('date_from')) or 'Desde el primer registro'} a {filters.get('date_to', metadata.get('date_to')) or 'Hasta la generación'}",
        "Filtros": json.dumps(filters, ensure_ascii=False, default=str) if filters else "Sin filtros adicionales",
        "Registros": len(rows),
    }


def _counts(rows):
    if any(row.get("record_type") == "evaluations" for row in rows):
        rows = [row for row in rows if row.get("record_type") == "evaluations"]
    return Counter(row.get("risk", row.get("severity", "sin_clasificar")) for row in rows)


def _chart_pdf(rows):
    counts = _counts(rows)
    drawing = Drawing(600, 115)
    maximum = max([counts.get(k, 0) for k in RISK_LABELS] + [1])
    for i, (key, label) in enumerate(RISK_LABELS.items()):
        y = 90 - i * 25
        drawing.add(String(0, y, label, fontName="Vigia", fontSize=9))
        if counts.get(key, 0):
            drawing.add(Rect(65, y - 3, 350 * counts.get(key, 0) / maximum, 12, fillColor=colors.HexColor({"low": "#128674", "moderate": "#d39922", "high": "#dc6b38", "critical": "#c83048"}[key]), strokeColor=None))
        drawing.add(String(425, y, str(counts.get(key, 0)), fontName="Vigia", fontSize=9))
    return drawing


def _render_pdf(kind, rows, metadata):
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=1.5*cm, leftMargin=1.5*cm, topMargin=1.5*cm, bottomMargin=1.5*cm, title=TITLES[kind], author="Vigía Financiera")
    styles = getSampleStyleSheet()
    for name in ("Normal", "BodyText", "Title", "Heading2"):
        styles[name].fontName = "VigiaBold" if name in {"Title", "Heading2"} else "Vigia"
    styles.add(ParagraphStyle(name="SmallCell", fontName="Vigia", fontSize=8, leading=11, splitLongWords=True, alignment=TA_LEFT))
    story = [Paragraph(escape(TITLES[kind]), styles["Title"]), Paragraph("Evidencia operacional para revisar indicadores de riesgo de phishing financiero.", styles["Normal"]), Spacer(1, 12)]
    for key, value in _metadata(kind, rows, metadata).items():
        story.append(Paragraph(f"<b>{key}:</b> {escape(_text(value))}", styles["Normal"]))
    story += [Spacer(1, 14), Paragraph("Resumen estadístico", styles["Heading2"])]
    if kind == "consolidated":
        counts = Counter(row.get("record_type", "evaluations") for row in rows)
        story.append(Paragraph(f"Evaluaciones: {counts['evaluations']} · Alertas: {counts['alerts']} · Incidentes: {counts['incidents']}. La distribución de riesgo representa las evaluaciones, sin volver a contar sus alertas e incidentes.", styles["Normal"]))
    story.append(_chart_pdf(rows))
    columns = COLUMNS.get(kind, COLUMNS["evaluations"])
    table_data = [[Paragraph(label, styles["SmallCell"]) for _, label in columns]]
    for row in rows:
        table_data.append([Paragraph(escape(_text(v)), styles["SmallCell"]) for v in _values(row, columns)])
    if not rows:
        story.append(Paragraph("No hay registros para los filtros seleccionados.", styles["Normal"]))
    else:
        width = landscape(A4)[0] - 3 * cm
        proportions = [.16, .25, .09, .1, .15, .25] if kind == "alerts" else [.10, .16, .21, .13, .14, .26] if kind == "consolidated" else [.17, .31, .11, .12, .13, .16]
        widths = [width * f for f in proportions]
        table = PDFTable(table_data, colWidths=widths, repeatRows=1, hAlign="LEFT", splitInRow=1)
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dceeea")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f5f7")]), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
        story.append(table)
    story += [Spacer(1, 12), Paragraph("La puntuación expresa indicadores de riesgo. No acredita un hackeo ni garantiza que un sitio sea seguro. Las evidencias externas pueden estar incompletas.", styles["Normal"])]

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Vigia", 8)
        canvas.drawString(1.5*cm, cm*.75, "Vigía Financiera | Uso académico | Datos sujetos a control de acceso")
        canvas.drawRightString(landscape(A4)[0]-1.5*cm, cm*.75, f"Página {document.page}")
        canvas.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def _chart_png(rows):
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new("RGB", (1000, 240), "white")
    draw = ImageDraw.Draw(image)
    counts = _counts(rows)
    maximum = max([counts.get(k, 0) for k in RISK_LABELS] + [1])
    font = ImageFont.truetype(str(FONTS / "DejaVuSans.ttf"), 22)
    for i, (key, label) in enumerate(RISK_LABELS.items()):
        y = 15 + 55*i
        draw.text((10, y+8), label, fill="#172c31", font=font)
        if counts.get(key, 0):
            draw.rectangle((170, y, 170+600*counts.get(key, 0)/maximum, y+30), fill={"low": "#128674", "moderate": "#d39922", "high": "#dc6b38", "critical": "#c83048"}[key])
        draw.text((810, y+4), str(counts.get(key, 0)), fill="#172c31", font=font)
    output = BytesIO()
    image.save(output, format="PNG")
    output.seek(0)
    return output


def _render_docx(kind, rows, metadata):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(11.69), Inches(8.27)
    section.top_margin = section.bottom_margin = Inches(.6)
    section.left_margin = section.right_margin = Inches(.65)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    document.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    for style in document.styles:
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    document.styles["Heading 1"].font.color.rgb = RGBColor(21, 94, 89)
    document.add_heading(TITLES[kind], 0)
    document.add_paragraph("Este reporte presenta indicadores de riesgo y registros operacionales del período seleccionado para apoyar la revisión autorizada.")
    for key, value in _metadata(kind, rows, metadata).items():
        paragraph = document.add_paragraph()
        paragraph.add_run(key+": ").bold = True
        paragraph.add_run(_text(value))
    document.add_heading("Resumen estadístico", 1)
    if kind == "consolidated":
        counts = Counter(row.get("record_type", "evaluations") for row in rows)
        document.add_paragraph(f"Evaluaciones: {counts['evaluations']} · Alertas: {counts['alerts']} · Incidentes: {counts['incidents']}. El gráfico cuenta solo evaluaciones.")
    document.add_picture(_chart_png(rows), width=Inches(7.3))
    document.add_heading("Detalle de eventos", 1)
    columns = COLUMNS.get(kind, COLUMNS["evaluations"])
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.autofit = False
    proportions = [.10, .16, .21, .13, .14, .26] if kind == "consolidated" else [.16, .25, .09, .1, .15, .25] if kind == "alerts" else [.17, .25, .1, .12, .13, .23]
    available = section.page_width - section.left_margin - section.right_margin
    for column, proportion in zip(table.columns, proportions):
        column.width = int(available * proportion)
    for cell, (_, label) in zip(table.rows[0].cells, columns):
        cell.text = label
    header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    for row in rows:
        for cell, value in zip(table.add_row().cells, _values(row, columns)):
            cell.text = _text(value)
    for index, row in enumerate(table.rows):
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for cell, proportion in zip(row.cells, proportions):
            cell.width = int(available * proportion)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "155E59" if index == 0 else "F3F5F7" if index % 2 else "FFFFFF")
            cell._tc.get_or_add_tcPr().append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.line_spacing = 1
                if index == 0:
                    paragraph.paragraph_format.keep_with_next = True
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.font.bold = index == 0
                    run.font.color.rgb = RGBColor(255,255,255) if index == 0 else RGBColor(20,30,35)
    if not rows:
        document.add_paragraph("No hay registros para los filtros seleccionados.")
    document.add_paragraph("La puntuación expresa indicadores de riesgo; no acredita un hackeo ni garantiza que un sitio sea seguro.")
    footer = section.footer.paragraphs[0]
    footer.add_run("Vigía Financiera | Página ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.author = "Vigía Financiera"
    document.core_properties.title = TITLES[kind]
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _xlsx_value(value):
    # Prevent attacker controlled titles/domains from becoming Excel formulae.
    value = _text(value) if not isinstance(value, (int, float)) else value
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _render_xlsx(kind, rows, metadata):
    wb = Workbook()
    summary = wb.active
    summary.title = "Resumen"
    summary.append([TITLES[kind]])
    for key, value in _metadata(kind, rows, metadata).items():
        summary.append([key, _xlsx_value(value)])
    summary.append([])
    chart_start = summary.max_row + 1
    summary.append(["Nivel de riesgo", "Registros"])
    counts = _counts(rows)
    for key, label in RISK_LABELS.items():
        summary.append([label, counts.get(key, 0)])
    chart = BarChart()
    chart.title = "Distribución de riesgo"
    chart.y_axis.title = "Registros"
    chart.x_axis.title = "Nivel"
    chart.add_data(Reference(summary, min_col=2, min_row=chart_start, max_row=chart_start+4), titles_from_data=True)
    chart.set_categories(Reference(summary, min_col=1, min_row=chart_start+1, max_row=chart_start+4))
    summary.add_chart(chart, "D2")
    summary.column_dimensions["A"].width, summary.column_dimensions["B"].width = 38, 65
    detail = wb.create_sheet("Eventos")
    columns = COLUMNS.get(kind, COLUMNS["evaluations"])
    detail.append([label for _, label in columns])
    for row in rows:
        detail.append([_xlsx_value(v) for v in _values(row, columns)])
    if rows:
        table = Table(displayName="EventosVigia", ref=f"A1:F{len(rows)+1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        detail.add_table(table)
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions
    detail.print_title_rows = "1:1"
    detail.sheet_properties.pageSetUpPr.fitToPage = True
    detail.page_setup.orientation = "landscape"
    detail.page_setup.fitToWidth, detail.page_setup.fitToHeight = 1, 0
    for index, width in enumerate([25, 55, 20, 24, 28, 35], 1):
        from openpyxl.utils import get_column_letter
        detail.column_dimensions[get_column_letter(index)].width = width
    for sheet in (summary, detail):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="155E59")
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def render_report(kind: str, format: str, rows: list[dict], metadata: dict) -> bytes:
    if kind not in TITLES or format not in {"pdf", "docx", "xlsx"}:
        raise ValueError("Tipo o formato de reporte no válido.")
    if kind == "consolidated":
        rows = [{**row, "event_title": row.get("domain") or row.get("title", "Evento"),
                 "status": row.get("status") or ("read" if row.get("read") else "unread") if row.get("record_type") == "alerts" else row.get("status"),
                 "event_risk": f"{row.get('score', '')} {RISK_LABELS.get(row.get('risk', row.get('severity')), '')}".strip(),
                 "event_detail": row.get("explanation") or row.get("message") or row.get("timeline", "")}
                for row in rows]
    return {"pdf": _render_pdf, "docx": _render_docx, "xlsx": _render_xlsx}[format](kind, rows, metadata)
